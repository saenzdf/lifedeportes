#!/usr/bin/env python3
"""
Auditoría local lista (Excel) vs PDF imprimible: alineada con Kapso `print_qc_webhook_odoo.js`.
Por defecto compara multiconjunto por nombre+dorsal (`dorsal`); modo `full` = nombre+talla+dorsal.
PDF: capa de texto + OCR Tesseract (extract_print_pdf_local) o visión (extract_print_pdf_v3).

Uso:
  cd lifedeportes/kapso && . .venv-pdf/bin/activate
  pip install -r scripts/requirements-print-pdf.txt
  python scripts/compare_list_pdf_local.py --task 1748
  python scripts/compare_list_pdf_local.py --task 1748 --pdf-source vision
  # Sin Odoo: mismos archivos en disco (p. ej. lifedeportes/memory/)
  python scripts/compare_list_pdf_local.py --list-file ../memory/lista.xlsx --print-pdf ../memory/arte.pdf --pdf-source vision
"""

from __future__ import annotations

import argparse
import base64
import importlib.util
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path


def _load_epdf():
    ep = Path(__file__).resolve().parent / "extract_print_pdf_local.py"
    spec = importlib.util.spec_from_file_location("extract_print_pdf_local", ep)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    return mod


def _load_v3():
    ep = Path(__file__).resolve().parent / "extract_print_pdf_v3.py"
    spec = importlib.util.spec_from_file_location("extract_print_pdf_v3_cmp", ep)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    return mod


def normalize_dict_rows_like_pdf(rows: list[dict]) -> list[dict[str, str]]:
    """Misma normalización que tuples_to_pdf_rows (Excel vs PDF)."""
    out: list[dict[str, str]] = []
    for r in rows:
        n = str(r.get("nombre_uniforme") or "").strip()
        if not n:
            continue
        out.append(
            {
                "nombre_uniforme": n,
                "talla": normalize_talla_py(r.get("talla")),
                "numero": normalize_numero_py(r.get("numero")),
            }
        )
    return out


def build_alarms(
    cmp_result: dict,
    excel_n: int,
    pdf_n: int,
    *,
    compare_mode: str,
    talla_mismatches: list[dict[str, str]] | None = None,
) -> list[str]:
    alarms: list[str] = []
    if excel_n != pdf_n:
        alarms.append(f"Conteo distinto: excel={excel_n} filas, pdf={pdf_n} filas")
    if not cmp_result.get("ok"):
        miss = cmp_result.get("missing") or []
        extra = cmp_result.get("extra") or []
        kind = "par nombre+dorsal" if compare_mode == "dorsal" else "trío nombre+talla+dorsal"
        if miss:
            alarms.append(
                f"En Excel pero no en PDF ({kind}): {len(miss)} — muestra: "
                + "; ".join(
                    f"{x[0]} #{x[2]}" if len(x) >= 3 and compare_mode == "dorsal" else "|".join(x)
                    for x in miss[:5]
                )
                + (" …" if len(miss) > 5 else "")
            )
        if extra:
            alarms.append(
                f"En PDF pero no en Excel ({kind}): {len(extra)} — muestra: "
                + "; ".join(
                    f"{x[0]} #{x[2]}" if len(x) >= 3 and compare_mode == "dorsal" else "|".join(x)
                    for x in extra[:5]
                )
                + (" …" if len(extra) > 5 else "")
            )
    if cmp_result.get("pdf_text_empty"):
        alarms.append("PDF no produjo filas con nombre mientras Excel sí tiene datos")
    if talla_mismatches:
        n = len(talla_mismatches)
        sample = talla_mismatches[:4]
        parts = [f"{t.get('nombre_uniforme')} # {t.get('numero')}: Excel {t.get('talla_excel')} vs PDF {t.get('talla_pdf') or '—'}" for t in sample]
        alarms.append(f"Revisar talla (con mismo nombre+dorsal en ambos lados): {n} — " + "; ".join(parts) + (" …" if n > 4 else ""))
    return alarms


def strip_diacritics(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn")


def normalize_talla_py(value: object) -> str:
    s = str(value or "").strip()
    if not s:
        return ""
    s = strip_diacritics(s).strip()
    s = re.sub(r"^\s*tallas?\s*[.:]?\s*", "", s, flags=re.I).strip()
    s = re.sub(r"\s+", " ", s).strip()
    if not s:
        return ""
    mf = re.match(r"^(\d+)\.0$", s)
    if mf:
        return str(int(mf.group(1)))
    mc = re.match(r"^(\d+),0$", s)
    if mc:
        return str(int(mc.group(1)))
    if re.match(r"^\d+$", s):
        return str(int(s))
    if re.match(r"^(XXS|XS|S|M|L|XL|XXL|2XL|3XL|4XL)$", s, re.I):
        return s.upper()
    if re.match(r"^\d+\s*-\s*\d+$", s):
        return re.sub(r"\s+", "", s).upper()
    return s.upper()


def normalize_token_py(value: object) -> str:
    s = strip_diacritics(str(value or "").strip().upper())
    s = re.sub(r"\s+", " ", s)
    for a, b in zip("ÁÉÍÓÚÑ", "AEIOUN"):
        s = s.replace(a, b)
    return s


def normalize_nombre_uniforme_for_qc_key(nombre: object, *, name_normalize: bool) -> str:
    """
    Normalización extra solo para claves de comparación (no altera el texto mostrado).
    Quitar punto final tras inicial de una letra: "SOFIA M." -> "SOFIA M".
    """
    s = normalize_token_py(nombre)
    if not name_normalize:
        return s
    return re.sub(r"(\s+[A-Z])\.\s*$", r"\1", s)


def _levenshtein_1(a: str, b: str) -> bool:
    if a == b:
        return True
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(1 for i in range(len(a)) if a[i] != b[i]) <= 1
    # insert/delete
    shorter, longer = (a, b) if len(a) < len(b) else (b, a)
    for i in range(len(longer)):
        if longer[:i] + longer[i + 1 :] == shorter:
            return True
    return False


def resolve_equivalence_token(
    token: str, numero: str, rules: list[dict]
) -> str:
    """Mapea token a canónico por grupos en JSON (mismo min lex entre equivalentes)."""
    n = normalize_numero_py(numero)
    for rule in rules:
        want = rule.get("numero")
        if want is not None and str(want).strip() and normalize_numero_py(want) != n:
            continue
        toks = rule.get("equivalent_tokens")
        if not isinstance(toks, list):
            continue
        tset = {normalize_nombre_uniforme_for_qc_key(t, name_normalize=True) for t in toks}
        tset = {t for t in tset if t}
        if token in tset and tset:
            return min(tset)
    return token


def load_equivalence_rules_from_path(path: Path | None) -> list[dict]:
    if not path or not path.is_file():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    rules = data.get("rules")
    return list(rules) if isinstance(rules, list) else []


def name_normalize_enabled_from_env() -> bool:
    v = os.environ.get("PRINT_QC_NAME_NORMALIZE", "1").strip().lower()
    return v not in ("0", "false", "off", "no")


def equivalence_path_from_env() -> Path | None:
    raw = os.environ.get("PRINT_QC_NAME_EQUIVALENCE_PATH", "").strip()
    return Path(raw) if raw else None


def normalize_numero_py(value: object) -> str:
    s = str(value or "").strip().replace("#", "")
    if re.match(r"^\d+$", s):
        return str(int(s))
    return normalize_token_py(s)


def row_key_py(n: str, t: str, num: str) -> str:
    return f"{normalize_token_py(n)}|{normalize_token_py(t)}|{normalize_numero_py(num)}"


def row_key_dorsal_py(
    n: str,
    num: str,
    *,
    name_normalize: bool = True,
    equivalence_rules: list[dict] | None = None,
) -> str:
    """Clave nombre+dorsal (talla excluida). Dos jugadores distintos con mismo dorsal → claves distintas por nombre."""
    tok = normalize_nombre_uniforme_for_qc_key(n, name_normalize=name_normalize)
    if equivalence_rules:
        tok = resolve_equivalence_token(tok, num, equivalence_rules)
    return f"{tok}|{normalize_numero_py(num)}"


def split_dorsal_key(key: str) -> tuple[str, str]:
    i = key.rfind("|")
    if i < 0:
        return key, ""
    return key[:i], key[i + 1 :]


def compare_multisets_py(
    excel_rows: list[dict[str, str]], pdf_rows: list[dict[str, str]]
) -> dict:
    ex = [
        row_key_py(r["nombre_uniforme"], r["talla"], r["numero"])
        for r in excel_rows
        if r.get("nombre_uniforme")
    ]
    pr = [
        row_key_py(r["nombre_uniforme"], r["talla"], r["numero"])
        for r in pdf_rows
        if r.get("nombre_uniforme")
    ]
    c_ex = Counter(ex)
    c_pr = Counter(pr)
    missing: list[list[str]] = []
    extra: list[list[str]] = []
    for k, n in c_ex.items():
        for _ in range(n - c_pr.get(k, 0)):
            missing.append(k.split("|"))
    for k, n in c_pr.items():
        for _ in range(n - c_ex.get(k, 0)):
            extra.append(k.split("|"))
    ok = len(missing) == 0 and len(extra) == 0
    return {
        "ok": ok,
        "mode": "full",
        "excel_count": len(ex),
        "pdf_count": len(pr),
        "missing": missing,
        "extra": extra,
        "pdf_text_empty": len(pr) == 0 and len(ex) > 0,
    }


def compare_multisets_dorsal_py(
    excel_rows: list[dict[str, str]],
    pdf_rows: list[dict[str, str]],
    *,
    name_normalize: bool = True,
    equivalence_rules: list[dict] | None = None,
) -> dict:
    """Multiconjunto por (nombre, dorsal). Adecuado cuando la talla en PDF por IA/OCR es inestable."""
    kw = {"name_normalize": name_normalize, "equivalence_rules": equivalence_rules or []}
    ex = [
        row_key_dorsal_py(r["nombre_uniforme"], r["numero"], **kw)
        for r in excel_rows
        if r.get("nombre_uniforme")
    ]
    pr = [
        row_key_dorsal_py(r["nombre_uniforme"], r["numero"], **kw)
        for r in pdf_rows
        if r.get("nombre_uniforme")
    ]
    c_ex = Counter(ex)
    c_pr = Counter(pr)
    missing: list[list[str]] = []
    extra: list[list[str]] = []
    for k, n in c_ex.items():
        for _ in range(n - c_pr.get(k, 0)):
            nk, num = split_dorsal_key(k)
            missing.append([nk, "", num])
    for k, n in c_pr.items():
        for _ in range(n - c_ex.get(k, 0)):
            nk, num = split_dorsal_key(k)
            extra.append([nk, "", num])
    ok = len(missing) == 0 and len(extra) == 0
    return {
        "ok": ok,
        "mode": "dorsal",
        "excel_count": len(ex),
        "pdf_count": len(pr),
        "missing": missing,
        "extra": extra,
        "pdf_text_empty": len(pr) == 0 and len(ex) > 0,
    }


def talla_mismatches_py(
    excel_rows: list[dict[str, str]],
    pdf_rows: list[dict[str, str]],
    *,
    name_normalize: bool = True,
    equivalence_rules: list[dict] | None = None,
) -> list[dict[str, str]]:
    """Tras coincidencia dorsal, lista discrepancias de talla (parejas 1:1 por clave dorsal)."""
    kw = {"name_normalize": name_normalize, "equivalence_rules": equivalence_rules or []}

    def group(rows: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
        g: dict[str, list[dict[str, str]]] = defaultdict(list)
        for r in rows:
            if not r.get("nombre_uniforme"):
                continue
            k = row_key_dorsal_py(r["nombre_uniforme"], r["numero"], **kw)
            g[k].append(r)
        return g

    ge, gp = group(excel_rows), group(pdf_rows)
    out: list[dict[str, str]] = []
    for k, elist in ge.items():
        plist = gp.get(k) or []
        if len(elist) != len(plist):
            continue
        elist_s = sorted(elist, key=lambda r: normalize_token_py(r["nombre_uniforme"]))
        plist_s = sorted(plist, key=lambda r: normalize_token_py(r["nombre_uniforme"]))
        for e, p in zip(elist_s, plist_s):
            te = normalize_talla_py(e.get("talla"))
            tp = normalize_talla_py(p.get("talla"))
            name = str(e.get("nombre_uniforme") or "").strip()
            num = normalize_numero_py(e.get("numero"))
            if te and tp and te != tp:
                out.append(
                    {
                        "nombre_uniforme": name,
                        "numero": num,
                        "talla_excel": te,
                        "talla_pdf": tp,
                        "kind": "mismatch",
                    }
                )
            elif te and not tp:
                out.append(
                    {
                        "nombre_uniforme": name,
                        "numero": num,
                        "talla_excel": te,
                        "talla_pdf": "",
                        "kind": "talla_missing_pdf",
                    }
                )
    return out


def run_compare_mode(
    excel_rows: list[dict[str, str]],
    pdf_rows: list[dict[str, str]],
    *,
    mode: str,
    name_normalize: bool = True,
    equivalence_rules: list[dict] | None = None,
) -> tuple[dict, list[dict[str, str]]]:
    if mode == "full":
        return compare_multisets_py(excel_rows, pdf_rows), []
    cmp_d = compare_multisets_dorsal_py(
        excel_rows,
        pdf_rows,
        name_normalize=name_normalize,
        equivalence_rules=equivalence_rules,
    )
    talla: list[dict[str, str]] = (
        talla_mismatches_py(
            excel_rows,
            pdf_rows,
            name_normalize=name_normalize,
            equivalence_rules=equivalence_rules,
        )
        if cmp_d["ok"]
        else []
    )
    return cmp_d, talla


def qc_label(cmp_ok: bool, compare_mode: str, talla_mismatches: list[dict[str, str]]) -> str:
    if not cmp_ok:
        return "diff"
    if compare_mode == "dorsal" and talla_mismatches:
        return "warn"
    return "ok"


def heuristic_dorsal_diff_labels(compare: dict, *, name_normalize: bool) -> dict[str, list]:
    """Etiquetas para feedback humano: likely_punctuation, likely_ocr, review (solo mode dorsal)."""
    if compare.get("mode") != "dorsal":
        return {}
    missing = compare.get("missing") or []
    extra = compare.get("extra") or []

    def label_pair(a_nk: str, b_nk: str) -> str:
        ca = normalize_nombre_uniforme_for_qc_key(a_nk, name_normalize=True)
        cb = normalize_nombre_uniforme_for_qc_key(b_nk, name_normalize=True)
        if ca == cb:
            return "likely_punctuation"
        if _levenshtein_1(a_nk, b_nk):
            return "likely_ocr"
        return "review"

    miss_details = []
    for m in missing:
        num = m[2] if len(m) > 2 else ""
        cands = [e for e in extra if len(e) > 2 and e[2] == num]
        label = "review"
        paired = None
        if len(cands) == 1:
            paired = cands[0]
            label = label_pair(m[0], paired[0])
        elif len(cands) > 1:
            paired = cands[0]
            label = "review"
        miss_details.append({"row": m, "label": label, "paired_extra": paired})

    extra_details = []
    for e in extra:
        num = e[2] if len(e) > 2 else ""
        cands = [m for m in missing if len(m) > 2 and m[2] == num]
        label = "review"
        paired = None
        if len(cands) == 1:
            paired = cands[0]
            label = label_pair(paired[0], e[0])
        elif len(cands) > 1:
            paired = cands[0]
            label = "review"
        extra_details.append({"row": e, "label": label, "paired_missing": paired})

    return {"missing_details": miss_details, "extra_details": extra_details}


def norm_cell_header(v: object) -> str:
    return re.sub(r"\s+", " ", strip_diacritics(str(v or "")).strip().upper())


def looks_like_nombre_header(t: str) -> bool:
    if not t:
        return False
    if "NOMBRE" in t and "UNIFORME" in t:
        return True
    if "NOMBRE" in t and "QUE" in t and "LLEVAR" in t:
        return True
    if "NOMBRE" in t and ("LISTADO" in t or "JUGADOR" in t or "ALUMNO" in t):
        return True
    if t in ("NOMBRE EN UNIFORME", "NOMBRE"):
        return True
    if "LISTADO" in t and "NOMBRE" in t:
        return True
    return False


def looks_like_talla_header(t: str) -> bool:
    if not t:
        return False
    if "TALLA" in t:
        return True
    if any(x in t for x in ("CAMISETA", "CAMISA", "POLO", "BUSO")):
        return True
    return False


def looks_like_numero_header(t: str) -> bool:
    if not t:
        return False
    if "NUMERO" in t or "NÚMERO" in t:
        return True
    if "DORSAL" in t:
        return True
    if t == "#" or "N°" in t or "NO." in t or t.startswith("NO "):
        return True
    return False


def scan_header_indexes(grid: list[list[object]]) -> dict:
    max_row = min(80, len(grid))
    nombre_cols: list[tuple[int, int]] = []
    talla_cols: list[tuple[int, int]] = []
    numero_cols: list[tuple[int, int]] = []

    for r in range(max_row):
        row = grid[r] if r < len(grid) else []
        max_col = min(40, len(row))
        for c in range(max_col):
            t = norm_cell_header(row[c])
            if not t:
                continue
            if looks_like_nombre_header(t):
                nombre_cols.append((r, c))
            if looks_like_talla_header(t):
                talla_cols.append((r, c))
            if looks_like_numero_header(t):
                numero_cols.append((r, c))

    cn = ct = cnum = None
    header_row_idx = None

    strong_nombre = None
    for r, c in nombre_cols:
        t = norm_cell_header(grid[r][c])
        if "NOMBRE" in t and "UNIFORME" in t:
            strong_nombre = (r, c)
            break
    if strong_nombre:
        cn, header_row_idx = strong_nombre[1], strong_nombre[0]
    elif nombre_cols:
        cn, header_row_idx = nombre_cols[0][1], nombre_cols[0][0]

    talla_near = None
    if cn is not None:
        for x in talla_cols:
            if x[0] == header_row_idx or abs(x[1] - cn) <= 6:
                talla_near = x
                break
    elif talla_cols:
        talla_near = talla_cols[0]

    numero_near = None
    if cn is not None:
        for x in numero_cols:
            if x[0] == header_row_idx or abs(x[1] - cn) <= 8:
                numero_near = x
                break
    elif numero_cols:
        numero_near = numero_cols[0]

    if talla_near:
        ct = talla_near[1]
    if numero_near:
        cnum = numero_near[1]

    if cn is None and ct is None and cnum is None:
        return {"cn": 1, "ct": 2, "cnum": 3, "startIdx": 4}

    if cn is not None:
        if ct is None:
            ct = cn + 1
        if cnum is None:
            cnum = cn + 2

    start_idx = min(len(grid) - 1, header_row_idx + 1) if header_row_idx is not None else 4
    return {"cn": cn, "ct": ct, "cnum": cnum, "startIdx": start_idx}


def cell_scalar(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and abs(v - round(v)) < 1e-9:
        return str(int(round(v)))
    return str(v).strip()


def extract_rows_from_life_grid(grid: list[list[object]]) -> list[dict[str, str]]:
    picked = scan_header_indexes(grid)
    cn, ct, cnum, start_idx = picked["cn"], picked["ct"], picked["cnum"], picked["startIdx"]
    if cn is None or ct is None or cnum is None:
        cn, ct, cnum, start_idx = 1, 2, 3, 4

    out: list[dict[str, str]] = []
    skip_labels = {"NOMBRE EN UNIFORME", "NOMBRE", "TOTAL", "SUBTOTAL"}
    blank = 0
    for r in range(start_idx, len(grid)):
        row = grid[r] if r < len(grid) else []
        extend = max(cn, ct, cnum) + 1
        while len(row) < extend:
            row.append("")
        nombre = row[cn]
        if nombre is None or str(nombre).strip() == "":
            blank += 1
            if blank >= 25:
                break
            continue
        blank = 0
        ns = str(nombre).strip()
        ns_up = strip_diacritics(ns).upper()
        if ns_up in skip_labels:
            continue
        if "TOTAL UNIFORMES" in ns_up:
            continue
        out.append(
            {
                "nombre_uniforme": ns,
                "talla": normalize_talla_py(row[ct]),
                "numero": normalize_numero_py(cell_scalar(row[cnum])),
            }
        )
    return out


def xlsx_bytes_to_grid(data: bytes) -> tuple[list[list[object]], str]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), data_only=True)
    sheet_name = wb.sheetnames[0]
    low_pref = [
        n for n in wb.sheetnames if "formato" in n.lower() and "life" in n.lower()
    ]
    if low_pref:
        sheet_name = low_pref[0]
    ws = wb[sheet_name]
    grid: list[list[object]] = []
    for row in ws.iter_rows(values_only=True):
        grid.append(list(row))
    return grid, sheet_name


def parse_json_list(data: bytes) -> list[dict[str, str]]:
    payload = json.loads(data.decode("utf-8"))
    lines = payload if isinstance(payload, list) else payload.get("lines")
    if not isinstance(lines, list):
        return []
    out = []
    for line in lines:
        if not isinstance(line, dict):
            continue
        nombre = (
            line.get("nombre_uniforme")
            or line.get("nombre")
            or line.get("name")
            or line.get("NOMBRE EN UNIFORME")
        )
        if not nombre:
            continue
        talla_raw = (
            line.get("talla")
            or line.get("TALLA")
            or line.get("Talla")
            or line.get("size")
            or ""
        )
        num_raw = (
            line.get("numero") or line.get("NUMERO") or line.get("numero_uniforme") or ""
        )
        out.append(
            {
                "nombre_uniforme": str(nombre).strip(),
                "talla": normalize_talla_py(talla_raw),
                "numero": normalize_numero_py(num_raw),
            }
        )
    return out


def load_list_rows_from_bytes(name: str, data: bytes) -> list[dict[str, str]]:
    low = name.lower()
    if low.endswith(".json"):
        return parse_json_list(data)
    if low.endswith((".xlsx", ".xlsm", ".xltx")):
        grid, _sn = xlsx_bytes_to_grid(data)
        return extract_rows_from_life_grid(grid)
    raise ValueError(f"Formato de lista no soportado: {name}")


def pick_list_attachment(
    base: str, db: str, uid: int, pwd: str, task_id: int, sale_order_id: int | None
) -> tuple[int | None, str | None]:
    domain = [
        ["res_model", "=", "project.task"],
        ["res_id", "=", task_id],
        "|",
        ["name", "ilike", ".xlsx"],
        ["name", "ilike", ".json"],
    ]
    atts = odoo_jsonrpc(
        base,
        "object",
        "execute_kw",
        [
            db,
            uid,
            pwd,
            "ir.attachment",
            "search_read",
            [domain],
            {"fields": ["id", "name", "create_date"], "order": "create_date desc, id desc", "limit": 40},
        ],
    )
    lst = list(atts or [])
    json_first = next((a for a in lst if str(a.get("name") or "").lower().endswith(".json")), None)
    if json_first:
        return json_first["id"], json_first.get("name")
    xlsx = [a for a in lst if str(a.get("name") or "").lower().endswith(".xlsx")]
    pref = [a for a in xlsx if "formato pedido life" in str(a.get("name") or "").lower()]
    if pref:
        pref.sort(key=lambda a: str(a.get("create_date") or ""), reverse=True)
        return pref[0]["id"], pref[0].get("name")
    if xlsx:
        xlsx.sort(key=lambda a: str(a.get("create_date") or ""), reverse=True)
        return xlsx[0]["id"], xlsx[0].get("name")

    if not sale_order_id:
        return None, None
    sod = [
        ["res_model", "=", "sale.order"],
        ["res_id", "=", sale_order_id],
        ["name", "ilike", ".xlsx"],
    ]
    sox = odoo_jsonrpc(
        base,
        "object",
        "execute_kw",
        [
            db,
            uid,
            pwd,
            "ir.attachment",
            "search_read",
            [sod],
            {"fields": ["id", "name", "create_date"], "order": "create_date desc, id desc", "limit": 8},
        ],
    )
    so_list = list(sox or [])
    so_pref = [a for a in so_list if "formato pedido life" in str(a.get("name") or "").lower()]
    if so_pref:
        return so_pref[0]["id"], so_pref[0].get("name")
    if so_list:
        return so_list[0]["id"], so_list[0].get("name")
    return None, None


# Import Odoo helper + PDF helpers from sibling module
_epdf = _load_epdf()
load_dotenv = _epdf.load_dotenv
odoo_jsonrpc = _epdf.odoo_jsonrpc
pick_print_pdf = _epdf.pick_print_pdf
extract_text_fitz = _epdf.extract_text_fitz
extract_text_ocr = _epdf.extract_text_ocr
parse_three_token_lines = _epdf.parse_three_token_lines
parse_ocr_rows_v2 = _epdf.parse_ocr_rows_v2
merge_row_lists = _epdf.merge_row_lists


def tuples_to_pdf_rows(tuples: list[tuple[str, str, str]]) -> list[dict[str, str]]:
    return [
        {
            "nombre_uniforme": a.strip(),
            "talla": normalize_talla_py(b),
            "numero": normalize_numero_py(c),
        }
        for a, b, c in tuples
    ]


def extract_pdf_rows_pipeline(
    pdf_bytes: bytes,
    *,
    zoom: float,
    lang: str,
    psm: int,
    max_side: int,
    ocr_engine: str,
    enhance: bool,
    parse_mode: str,
    tesseract: str | None,
) -> tuple[list[dict[str, str]], dict]:
    text, meta_txt = extract_text_fitz(pdf_bytes)
    ocr_used = not (text or "").strip()
    meta: dict = {"text_layer": meta_txt, "ocr": None}
    if ocr_used:
        text, meta_ocr = extract_text_ocr(
            pdf_bytes,
            zoom=zoom,
            lang=lang,
            tesseract_cmd=tesseract,
            psm=psm,
            max_side=max_side,
            ocr_engine=ocr_engine,
            enhance=enhance,
        )
        meta["ocr"] = meta_ocr

    v1 = parse_three_token_lines(text, ocr=ocr_used)
    if parse_mode == "v1":
        tuples_out = v1
    elif parse_mode == "v2":
        tuples_out = parse_ocr_rows_v2(text) if ocr_used else []
    else:
        if ocr_used:
            tuples_out = merge_row_lists(v1, parse_ocr_rows_v2(text))
        else:
            tuples_out = v1

    return tuples_to_pdf_rows(tuples_out), meta


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    load_dotenv(root / ".env")

    _vp_env = os.environ.get("PRINT_QC_VISION_PROVIDER", "").strip().lower()
    _default_vp = _vp_env if _vp_env in ("gemini", "openai") else "gemini"
    _cmp_env = os.environ.get("PRINT_QC_COMPARE_MODE", "dorsal").strip().lower()
    _default_cmp = _cmp_env if _cmp_env in ("full", "dorsal") else "dorsal"

    parser = argparse.ArgumentParser(description="Comparar lista (Excel/JSON) vs PDF imprimible (local).")
    parser.add_argument("--task", type=int, default=0, help="project.task id (o usa --list-file + --print-pdf)")
    parser.add_argument(
        "--list-file",
        type=Path,
        default=None,
        help="Excel/JSON en disco; requiere --print-pdf (no usa Odoo)",
    )
    parser.add_argument(
        "--print-pdf",
        type=Path,
        default=None,
        help="PDF de impresión en disco; requiere --list-file",
    )
    parser.add_argument("--list-id", type=int, default=0, help="Forzar ir.attachment de lista (solo con --task)")
    parser.add_argument("--pdf-id", type=int, default=0, help="Forzar ir.attachment PDF (solo con --task)")
    parser.add_argument(
        "--compare-mode",
        choices=("full", "dorsal"),
        default=_default_cmp,
        help="full=nombre+talla+dorsal; dorsal=nombre+dorsal (recomendado con visión; ver PRINT_QC_COMPARE_MODE)",
    )
    parser.add_argument(
        "--no-name-normalize",
        action="store_true",
        help="Desactiva normalización de punto tras inicial en clave dorsal (override PRINT_QC_NAME_NORMALIZE)",
    )
    parser.add_argument(
        "--equivalence-file",
        type=Path,
        default=None,
        help="JSON equivalent_tokens (override env PRINT_QC_NAME_EQUIVALENCE_PATH)",
    )
    parser.add_argument(
        "--pdf-source",
        choices=("ocr", "vision"),
        default="ocr",
        help="ocr=pipeline Tesseract (extract_print_pdf_local); vision=Gemini/OpenAI como extract_print_pdf_v3",
    )
    parser.add_argument(
        "--vision-provider",
        choices=("gemini", "openai"),
        default=_default_vp,
    )
    parser.add_argument(
        "--gemini-model",
        default=os.environ.get("PRINT_QC_GEMINI_MODEL", "gemini-2.5-flash"),
    )
    parser.add_argument("--vision-model", default=os.environ.get("PRINT_QC_VISION_MODEL", "gpt-4o-mini"))
    parser.add_argument(
        "--vision-base-url",
        default=os.environ.get("OPENAI_API_BASE", "https://api.openai.com/v1"),
    )
    parser.add_argument("--zoom", type=float, default=2.5)
    parser.add_argument("--lang", default="spa+eng")
    parser.add_argument("--psm", type=int, default=6)
    parser.add_argument("--max-side", type=int, default=4500)
    parser.add_argument("--ocr-engine", choices=("string", "lines"), default="string")
    parser.add_argument("--enhance", action="store_true")
    parser.add_argument("--parse", choices=("merge", "v1", "v2"), default="merge")
    parser.add_argument("--tesseract", default="")
    args = parser.parse_args()

    use_files = bool(args.list_file and args.print_pdf)
    if bool(args.list_file) ^ bool(args.print_pdf):
        sys.exit("Usa juntos --list-file y --print-pdf, u omitir ambos y usar --task")
    if use_files and args.task:
        sys.exit("No combines --task con --list-file/--print-pdf")
    if not use_files and not args.task:
        sys.exit("Indica --task ID o --list-file + --print-pdf")

    list_id: int | None = None
    pdf_id: int | None = None
    list_name: str | None = None
    pdf_name: str | None = None
    task_id_out: int | None = int(args.task) if args.task else None

    if use_files:
        lf = args.list_file.expanduser().resolve()
        pf = args.print_pdf.expanduser().resolve()
        if not lf.is_file():
            sys.exit(f"No existe lista: {lf}")
        if not pf.is_file():
            sys.exit(f"No existe PDF: {pf}")
        list_bytes = lf.read_bytes()
        pdf_bytes = pf.read_bytes()
        list_name = lf.name
        pdf_name = pf.name
    else:
        base = os.environ.get("ODOO_URL", "").strip()
        db = os.environ.get("ODOO_DB", "").strip()
        user = os.environ.get("ODOO_USERNAME", "").strip()
        pwd = os.environ.get("ODOO_PASSWORD", "").strip()
        if not all([base, db, user, pwd]):
            sys.exit("Faltan variables Odoo en .env")

        uid = odoo_jsonrpc(base, "common", "authenticate", [db, user, pwd, {}])
        if not uid:
            sys.exit("authenticate failed")

        tasks = odoo_jsonrpc(
            base,
            "object",
            "execute_kw",
            [
                db,
                uid,
                pwd,
                "project.task",
                "read",
                [[args.task]],
                {"fields": ["id", "sale_order_id"]},
            ],
        )
        task = tasks[0] if tasks else None
        so_id = None
        if task and task.get("sale_order_id"):
            so_id = task["sale_order_id"][0] if isinstance(task["sale_order_id"], list) else task["sale_order_id"]

        list_id = args.list_id or None
        if not list_id:
            list_id, list_name = pick_list_attachment(base, db, uid, pwd, args.task, so_id)
        if not list_id:
            sys.exit("No se encontró Excel/JSON de lista en la tarea ni en el pedido.")

        pdf_id = args.pdf_id or None
        if not pdf_id:
            domain = [
                ["res_model", "=", "project.task"],
                ["res_id", "=", args.task],
                ["mimetype", "=", "application/pdf"],
            ]
            pdfs = odoo_jsonrpc(
                base,
                "object",
                "execute_kw",
                [
                    db,
                    uid,
                    pwd,
                    "ir.attachment",
                    "search_read",
                    [domain],
                    {"fields": ["id", "name", "create_date"], "order": "create_date desc, id desc", "limit": 30},
                ],
            )
            picked = pick_print_pdf(list(pdfs or []))
            if not picked:
                sys.exit("No hay PDF en la tarea.")
            pdf_id = picked["id"]
            pdf_name = picked.get("name")

        lst_meta = odoo_jsonrpc(
            base,
            "object",
            "execute_kw",
            [db, uid, pwd, "ir.attachment", "read", [[list_id]], {"fields": ["name", "datas"]}],
        )
        pdf_meta = odoo_jsonrpc(
            base,
            "object",
            "execute_kw",
            [db, uid, pwd, "ir.attachment", "read", [[pdf_id]], {"fields": ["name", "datas"]}],
        )
        la = lst_meta[0]
        pa = pdf_meta[0]
        if not list_name:
            list_name = la.get("name")
        if not pdf_name:
            pdf_name = pa.get("name")
        if not la.get("datas") or not pa.get("datas"):
            sys.exit("Adjunto sin contenido")

        list_bytes = base64.standard_b64decode(la["datas"])
        pdf_bytes = base64.standard_b64decode(pa["datas"])

    excel_rows = load_list_rows_from_bytes(str(list_name or ""), list_bytes)

    if args.pdf_source == "vision":
        v3 = _load_v3()
        gemini_key = (
            os.environ.get("GEMINI_API_KEY", "").strip()
            or os.environ.get("GOOGLE_API_KEY", "").strip()
        )
        openai_key = os.environ.get("OPENAI_API_KEY", "").strip()
        png = v3.render_full_page_png(pdf_bytes, args.zoom, args.max_side)
        if args.vision_provider == "gemini":
            if not gemini_key:
                sys.exit("GEMINI_API_KEY o GOOGLE_API_KEY requerido para --pdf-source vision con gemini")
            raw_rows, vision_meta = v3.vision_gemini_extract(
                png, api_key=gemini_key, model=args.gemini_model.strip()
            )
        else:
            if not openai_key:
                sys.exit("OPENAI_API_KEY requerido para --pdf-source vision con openai")
            raw_rows, vision_meta = v3.vision_openai_extract(
                png,
                api_key=openai_key,
                model=args.vision_model,
                base_url=args.vision_base_url,
            )
        pdf_rows = normalize_dict_rows_like_pdf(raw_rows)
        pdf_pipeline_meta = {"source": "vision", "vision_provider": args.vision_provider, "vision": vision_meta}
        if isinstance(vision_meta, dict) and vision_meta.get("error"):
            sys.exit(f"Visión falló: {vision_meta.get('error')} — {vision_meta}")
    else:
        tess = args.tesseract.strip() or os.environ.get("TESSERACT_CMD", "").strip() or None
        pdf_rows, pdf_pipeline_meta = extract_pdf_rows_pipeline(
            pdf_bytes,
            zoom=args.zoom,
            lang=args.lang,
            psm=args.psm,
            max_side=args.max_side,
            ocr_engine=args.ocr_engine,
            enhance=args.enhance,
            parse_mode=args.parse,
            tesseract=tess,
        )

    eq_file = args.equivalence_file if args.equivalence_file else equivalence_path_from_env()
    equiv_rules = load_equivalence_rules_from_path(Path(eq_file) if eq_file else None)
    name_norm = name_normalize_enabled_from_env() and not args.no_name_normalize

    cmp_result, talla_mm = run_compare_mode(
        excel_rows,
        pdf_rows,
        mode=args.compare_mode,
        name_normalize=name_norm,
        equivalence_rules=equiv_rules,
    )
    qc = qc_label(bool(cmp_result.get("ok")), args.compare_mode, talla_mm)
    alarms = build_alarms(
        cmp_result,
        len(excel_rows),
        len(pdf_rows),
        compare_mode=args.compare_mode,
        talla_mismatches=talla_mm,
    )

    out: dict = {
        "task_id": task_id_out,
        "list_attachment_id": list_id,
        "list_name": list_name,
        "print_attachment_id": pdf_id,
        "pdf_name": pdf_name,
        "compare_mode": args.compare_mode,
        "name_qc": {
            "normalize_enabled": name_norm,
            "equivalence_rules_count": len(equiv_rules),
            "equivalence_path": str(eq_file) if eq_file else None,
        },
        "pdf_source": args.pdf_source,
        "counts": {"excel_rows": len(excel_rows), "pdf_rows": len(pdf_rows)},
        "qc": qc,
        "alarms": alarms,
        "compare": cmp_result,
        "talla_mismatches": talla_mm,
        "pdf_pipeline": pdf_pipeline_meta,
    }
    print(json.dumps(out, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
